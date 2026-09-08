"""
modules/portfolio/router.py —— 持仓盈亏管理API

从 api/portfolio.py 迁移（v2.5.5架构重构 - 第4步按领域模块重组）
负责：持仓列表、投资记录CRUD、OCR上传、Excel导入、实时盈亏更新
"""
import os
import csv
import io
import re
import time
import datetime
from typing import Optional
from fastapi import APIRouter, HTTPException, UploadFile, File
from pydantic import BaseModel
import db
from modules.logging.audit import audit_log, AuditAction
from modules.datasource.eastmoney import EastMoneySource
from security.crypto import encrypt, decrypt

router = APIRouter(prefix='/api/advanced', tags=['持仓盈亏管理'])


# ============================================================
# 性能优化：TTL内存缓存
# ============================================================
_portfolio_cache = {'data': None, 'timestamp': 0}
_CACHE_TTL_TRADING = 30       # 交易时段缓存30秒
_CACHE_TTL_NON_TRADING = 300  # 非交易时段缓存5分钟


def _is_trading_hours() -> bool:
    now = datetime.datetime.now()
    return now.weekday() < 5 and 9 <= now.hour < 15


def invalidate_portfolio_cache():
    _portfolio_cache['data'] = None
    _portfolio_cache['timestamp'] = 0


def get_db():
    return db.get_conn()


def _decrypt_record(record):
    if record and record.get("note"):
        record["note"] = decrypt(record["note"])
    return record


def _ensure_profit_history_table(db_conn):
    """确保持仓收益历史表存在"""
    db_conn.execute('''
        CREATE TABLE IF NOT EXISTS portfolio_profit_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL,
            total_amount REAL DEFAULT 0,
            total_profit REAL DEFAULT 0,
            total_cost REAL DEFAULT 0,
            profit_pct REAL DEFAULT 0,
            daily_profit REAL DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now', 'localtime')),
            UNIQUE(date)
        )
    ''')
    db_conn.commit()


def _snapshot_daily_profit(db_conn):
    """记录当日持仓收益快照（幂等，同一天只记录一次）"""
    _ensure_profit_history_table(db_conn)
    today = datetime.date.today().isoformat()

    # 检查今天是否已有快照
    existing = db_conn.execute(
        'SELECT id FROM portfolio_profit_history WHERE date = ?', (today,)
    ).fetchone()
    if existing:
        return

    portfolios = db_conn.execute('SELECT * FROM portfolio').fetchall()
    total_amount = sum(p['total_amount'] or 0 for p in portfolios)
    total_profit = sum(p['total_profit'] or 0 for p in portfolios)
    total_cost = sum((p['avg_cost'] or 0) * (p['total_shares'] or 0) for p in portfolios)
    profit_pct = round((total_profit / total_cost * 100) if total_cost > 0 else 0, 2)

    # 计算单日收益（与前一交易日对比）
    prev = db_conn.execute(
        'SELECT total_profit FROM portfolio_profit_history WHERE date < ? ORDER BY date DESC LIMIT 1',
        (today,)
    ).fetchone()
    daily_profit = round(total_profit - (prev['total_profit'] if prev else 0), 2)

    db_conn.execute('''
        INSERT OR IGNORE INTO portfolio_profit_history
        (date, total_amount, total_profit, total_cost, profit_pct, daily_profit)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (today, total_amount, total_profit, total_cost, profit_pct, daily_profit))
    db_conn.commit()


# ============================================================
# 数据模型
# ============================================================

class InvestmentRecordCreate(BaseModel):
    code: str
    name: Optional[str] = None
    type: str = 'buy'  # buy/sell/dividend
    date: str
    amount: float
    shares: Optional[float] = None
    nav: Optional[float] = None
    fee: float = 0
    profit_target: Optional[float] = None
    note: Optional[str] = None
    source: str = 'manual'


class InvestmentUpdate(BaseModel):
    code: Optional[str] = None
    name: Optional[str] = None
    type: Optional[str] = None
    date: Optional[str] = None
    amount: Optional[float] = None
    shares: Optional[float] = None
    nav: Optional[float] = None
    fee: Optional[float] = None
    profit_target: Optional[float] = None
    note: Optional[str] = None


# ============================================================
# 持仓管理API
# ============================================================

@router.get('/portfolio', summary='获取持仓列表', responses={200: {'description': '持仓列表及汇总（含实时盈亏）'}})
async def list_portfolio():
    """获取持仓列表（含实时盈亏，TTL缓存）"""
    # 检查缓存
    now_ts = time.time()
    ttl = _CACHE_TTL_TRADING if _is_trading_hours() else _CACHE_TTL_NON_TRADING
    if _portfolio_cache['data'] and now_ts - _portfolio_cache['timestamp'] < ttl:
        return _portfolio_cache['data']

    db_conn = get_db()
    _update_portfolio_realtime(db_conn)
    _snapshot_daily_profit(db_conn)
    rows = db_conn.execute(
        'SELECT * FROM portfolio ORDER BY total_amount DESC'
    ).fetchall()
    total_amount = sum(r['total_amount'] or 0 for r in rows)
    total_profit = sum(r['total_profit'] or 0 for r in rows)
    total_cost = sum((r['avg_cost'] or 0) * (r['total_shares'] or 0) for r in rows)

    # 批量查询买入记录数量（1次GROUP BY替代N次COUNT）
    buy_counts = {row['code']: row['buy_count'] for row in db_conn.execute(
        "SELECT code, COUNT(*) as buy_count FROM investment_records WHERE type = 'buy' GROUP BY code"
    ).fetchall()}

    portfolio_rows = []
    for r in rows:
        d = dict(r)
        d['buy_count'] = buy_counts.get(r['code'], 0)
        portfolio_rows.append(d)
    # v2.9.8: 设了止盈目标的买入记录（按笔）——止盈汇总弹窗数据源。
    # 修复：1.过滤无效记录(buy_nav=0或shares=0) 2.标记已卖出记录 3.收益率口径修正(分母=实际成本)
    # 先查询所有卖出记录，按code分组计算总卖出份额
    sell_rows = db_conn.execute('''
        SELECT code, SUM(shares) as total_sold
        FROM investment_records
        WHERE type = 'sell'
        GROUP BY code
    ''').fetchall()
    sell_map = {r['code']: (r['total_sold'] or 0) for r in sell_rows}

    # 查询当前还持有的基金（用于判断是否已全部卖出）
    holding_codes = {r['code'] for r in db_conn.execute(
        'SELECT code FROM portfolio WHERE total_shares > 0'
    ).fetchall()}

    profit_rows = db_conn.execute('''
        SELECT r.id, r.code, r.date, r.nav AS buy_nav, r.amount, r.shares, r.fee, r.profit_target,
               COALESCE(f.name, r.name, r.code) AS fund_name, COALESCE(f.nav, r.nav) AS current_nav
        FROM investment_records r
        LEFT JOIN funds f ON f.code = r.code
        WHERE r.type = 'buy' AND r.profit_target IS NOT NULL AND r.profit_target > 0
          AND r.nav > 0 AND r.shares > 0
        ORDER BY r.date DESC, r.id DESC
    ''').fetchall()
    profit_records = []
    for pr in profit_rows:
        d = dict(pr)
        cur = d.get('current_nav') or 0
        buy_nav = d.get('buy_nav') or 0
        shares = d.get('shares') or 0
        fee = d.get('fee') or 0
        target = d.get('profit_target') or 0

        # 判断是否已卖出：该基金不在持仓中，或总卖出份额 >= 该笔买入份额
        total_sold = sell_map.get(d['code'], 0)
        is_sold = (d['code'] not in holding_codes) or (total_sold >= shares)

        # v2.9.8: 已卖出的记录直接过滤，不显示在止盈弹窗中
        if is_sold:
            continue

        # 收益率计算：分母=实际成本(buy_nav * shares + fee)，分子=当前市值-成本-手续费
        pct = 0.0
        cost = buy_nav * shares + fee
        if cur > 0 and buy_nav > 0 and cost > 0:
            profit_val = (cur - buy_nav) * shares - fee
            pct = round(profit_val / cost * 100, 2)

        d['profit_pct'] = pct
        d['reached'] = pct >= target
        d['status'] = 'holding'
        d['remaining_shares'] = max(0, shares - total_sold)

        # v2.9.29 修复：止盈提醒只显示已达标的记录（收益率 >= 止盈目标）
        # 未达标的记录不在止盈提醒中显示，在持仓详情页管理止盈目标
        if not d['reached']:
            continue

        profit_records.append(d)
    result = {
        'ok': True,
        'portfolio': portfolio_rows,
        'profit_records': profit_records,
        'summary': {
            'total_amount': round(total_amount, 2),
            'total_profit': round(total_profit, 2),
            'total_cost': round(total_cost, 2),
            'profit_pct': round((total_profit / total_cost * 100) if total_cost > 0 else 0, 2),
            'fund_count': len(rows)
        }
    }
    # 更新缓存
    _portfolio_cache['data'] = result
    _portfolio_cache['timestamp'] = time.time()
    return result


@router.get('/portfolio/{code}', summary='获取单只基金持仓详情', responses={200: {'description': '持仓详情及投资记录'}, 404: {'description': '未找到该基金持仓'}})
async def get_portfolio_detail(code: str):
    """获取单只基金持仓详情"""
    db_conn = get_db()
    portfolio = db_conn.execute('SELECT * FROM portfolio WHERE code = ?', (code,)).fetchone()
    if not portfolio:
        raise HTTPException(404, '未找到该基金持仓')
    records = db_conn.execute('''
        SELECT * FROM investment_records
        WHERE code = ?
        ORDER BY date DESC, id DESC
    ''', (code,)).fetchall()
    return {'ok': True, 'portfolio': dict(portfolio),
            'records': [_decrypt_record(r) for r in _attach_record_status(records)]}


def _attach_record_status(records):
    """为投资记录附加持仓状态（FIFO 抵扣计算）。

    卖出记录按时间正序抵扣最早的未清仓买入记录：
    - buy:  remaining_shares=剩余份额, status=holding(持有中)/partial(部分卖出)/sold(已卖出)
    - sell: status=sold
    前端据此显示"持有中/已卖出"状态标签，买入记录全部卖出后左侧色条置灰。
    """
    recs = [dict(r) for r in records]
    buys = sorted([r for r in recs if r['type'] == 'buy'],
                  key=lambda r: (r['date'] or '', r['id']))
    sells = sorted([r for r in recs if r['type'] == 'sell'],
                   key=lambda r: (r['date'] or '', r['id']))
    remaining = {b['id']: (b['shares'] or 0) for b in buys}
    for s in sells:
        need = s['shares'] or 0
        for b in buys:
            if need <= 1e-4:
                break
            rid = b['id']
            if remaining.get(rid, 0) > 1e-4:
                deduct = min(remaining[rid], need)
                remaining[rid] -= deduct
                need -= deduct
    for d in recs:
        if d['type'] == 'buy':
            total = d['shares'] or 0
            rem = round(remaining.get(d['id'], 0), 2)
            d['remaining_shares'] = rem
            if total <= 0:
                d['status'] = 'holding'
            elif rem <= 1e-4:
                d['status'] = 'sold'
            elif rem < total - 1e-4:
                d['status'] = 'partial'
            else:
                d['status'] = 'holding'
        elif d['type'] == 'sell':
            d['status'] = 'sold'
        else:
            d['status'] = 'holding'
    return recs


class BatchProfitTarget(BaseModel):
    code: str
    profit_target: float


@router.put('/portfolio/investment/batch-target', summary='批量设置止盈目标',
            responses={200: {'description': '更新成功'}, 400: {'description': '参数校验失败'}})
async def batch_profit_target(data: BatchProfitTarget):
    """一次性更新该基金全部买入记录的止盈目标（原子单条 SQL）。

    注意：必须注册在 /portfolio/investment/{record_id} 之前，
    否则 "batch-target" 会先被 {record_id:int} 匹配并 422。
    此前端逐笔并发 PUT，多笔记录并发写 SQLite 偶发 database is locked，
    部分失败导致详情页预期收益(取max)残留旧值——表现为"调整止盈后没变化"。
    """
    db_conn = get_db()
    if not data.profit_target or data.profit_target <= 0:
        raise HTTPException(400, '止盈目标必须大于0')
    fund = db_conn.execute('SELECT code FROM funds WHERE code = ?', (data.code,)).fetchone()
    if not fund:
        raise HTTPException(404, f'基金代码 {data.code} 不存在')
    cur = db_conn.execute(
        "UPDATE investment_records SET profit_target = ? WHERE code = ? AND type = 'buy'",
        (data.profit_target, data.code))
    db_conn.commit()
    audit_log(AuditAction.PORTFOLIO_UPDATE, f"批量设置止盈目标: {data.code} target={data.profit_target} 更新{cur.rowcount}笔",
              extra={"code": data.code, "profit_target": data.profit_target})
    invalidate_portfolio_cache()
    return {'ok': True, 'updated': cur.rowcount}


@router.post('/portfolio/investment', summary='添加投资记录', responses={200: {'description': '添加成功'}, 400: {'description': '参数校验失败'}, 404: {'description': '基金不存在'}})
async def add_investment(data: InvestmentRecordCreate):
    """添加投资记录（买入/卖出/分红）"""
    db_conn = get_db()

    # ========== 参数校验 ==========
    valid_types = ['buy', 'sell', 'dividend', '买入', '卖出', '分红']
    if data.type not in valid_types:
        raise HTTPException(400, f'无效的投资类型: {data.type}，支持: buy/sell/dividend')

    type_map = {'买入': 'buy', '卖出': 'sell', '分红': 'dividend'}
    invest_type = type_map.get(data.type, data.type)

    fund_row = db_conn.execute('SELECT code, name, nav FROM funds WHERE code = ?', (data.code,)).fetchone()
    if not fund_row:
        raise HTTPException(404, f'基金代码 {data.code} 不存在，请先添加到自选或检查代码是否正确')

    try:
        from datetime import datetime as dt
        dt.strptime(data.date, '%Y-%m-%d')
    except ValueError:
        raise HTTPException(400, f'无效的日期格式: {data.date}，请使用 YYYY-MM-DD 格式')

    amount = data.amount or 0
    shares = data.shares or 0
    if invest_type in ['buy', 'sell'] and amount <= 0 and shares <= 0:
        raise HTTPException(400, '投资金额和份额不能同时为0，请至少填写一个')

    if invest_type == 'sell' and shares > 0:
        portfolio = db_conn.execute('SELECT total_shares FROM portfolio WHERE code = ?', (data.code,)).fetchone()
        current_shares = portfolio['total_shares'] if portfolio else 0
        if shares > current_shares:
            raise HTTPException(400, f'卖出份额({shares})超过当前持有份额({current_shares})')

    # ========== 净值查询（严格兜底） ==========
    nav = data.nav
    if not nav or nav <= 0:
        nav_row = db_conn.execute(
            'SELECT ljjz FROM nav_history WHERE code = ? AND date = ?',
            (data.code, data.date)
        ).fetchone()
        if nav_row and nav_row[0] and nav_row[0] > 0:
            nav = nav_row[0]

    if not nav or nav <= 0:
        nav_row = db_conn.execute(
            'SELECT ljjz FROM nav_history WHERE code = ? AND date <= ? ORDER BY date DESC LIMIT 1',
            (data.code, data.date)
        ).fetchone()
        if nav_row and nav_row[0] and nav_row[0] > 0:
            nav = nav_row[0]

    if not nav or nav <= 0:
        if fund_row and fund_row['nav'] and fund_row['nav'] > 0:
            nav = fund_row['nav']

    if not nav or nav <= 0:
        raise HTTPException(400, f'无法获取基金 {data.code} 的净值数据，请检查基金代码或稍后重试')

    # ========== 份额/金额自动计算 ==========
    if (not shares or shares <= 0) and amount > 0 and nav > 0:
        shares = round(amount / nav, 2)

    if (not amount or amount <= 0) and shares > 0 and nav > 0:
        amount = round(shares * nav, 2)

    if invest_type == 'sell':
        amount = abs(amount)
        shares = abs(shares)

    # ========== 基金名称兜底 ==========
    name = data.name or fund_row['name'] or data.code

    # ========== 插入投资记录 ==========
    cursor = db_conn.execute('''
        INSERT INTO investment_records
        (code, name, type, date, amount, shares, nav, fee, profit_target, note, source)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (data.code, name, invest_type, data.date, amount, shares, nav,
          data.fee or 0, data.profit_target, encrypt(data.note) if data.note else None, data.source or 'manual'))
    db_conn.commit()

    # ========== 重新计算持仓汇总 ==========
    _recalculate_portfolio(db_conn, data.code)
    audit_log(AuditAction.PORTFOLIO_ADD, f"添加投资记录: {data.code} {invest_type} 金额={amount} 份额={shares}", extra={"code": data.code, "type": invest_type, "amount": amount, "shares": shares})
    invalidate_portfolio_cache()

    return {
        'ok': True,
        'id': cursor.lastrowid,
        'shares': shares,
        'nav': nav,
        'name': name,
        'amount': amount,
        'type': invest_type
    }


@router.delete('/portfolio/investment/{record_id}', summary='删除投资记录', responses={200: {'description': '删除成功'}, 404: {'description': '记录不存在'}})
async def delete_investment(record_id: int):
    """删除投资记录"""
    db_conn = get_db()
    record = db_conn.execute('SELECT code FROM investment_records WHERE id = ?', (record_id,)).fetchone()
    if not record:
        raise HTTPException(404, '未找到该投资记录')
    code = record[0]
    db_conn.execute('DELETE FROM investment_records WHERE id = ?', (record_id,))
    db_conn.commit()
    _recalculate_portfolio(db_conn, code)
    audit_log(AuditAction.PORTFOLIO_REMOVE, f"删除投资记录: id={record_id} code={code}", extra={"record_id": record_id, "code": code})
    invalidate_portfolio_cache()
    return {'ok': True}


@router.delete('/portfolio/{code}', summary='删除整个持仓', responses={200: {'description': '删除成功'}})
async def delete_portfolio(code: str):
    """删除整个持仓（该基金的所有投资记录和持仓汇总）"""
    db_conn = get_db()
    db_conn.execute('DELETE FROM investment_records WHERE code = ?', (code,))
    db_conn.execute('DELETE FROM portfolio WHERE code = ?', (code,))
    # v2.8.1: 同步清理该基金的止盈提醒规则，避免悬空规则继续触发
    db_conn.execute("DELETE FROM alert_rules WHERE type = 'profit' AND code = ?", (code,))
    db_conn.commit()
    audit_log(AuditAction.PORTFOLIO_REMOVE, f"删除整个持仓: {code}", extra={"code": code})
    invalidate_portfolio_cache()
    return {'ok': True, 'deleted': code}


@router.delete('/portfolio/{code}/records', summary='清除该基金所有投资记录', responses={200: {'description': '清除成功'}})
async def clear_portfolio_records(code: str):
    """清除该基金的所有投资记录（买入+卖出），保留持仓本身（变为空仓）。

    与「删除整个持仓」的区别：
    - 删除持仓：连 portfolio 表记录一起删，持仓列表中不再显示
    - 清除记录：只删 investment_records，portfolio 保留为空仓，可继续添加记录
    """
    db_conn = get_db()
    # 检查持仓是否存在
    portfolio = db_conn.execute('SELECT code FROM portfolio WHERE code = ?', (code,)).fetchone()
    if not portfolio:
        raise HTTPException(404, '未找到该基金持仓')
    # 删除所有投资记录
    cursor = db_conn.execute('DELETE FROM investment_records WHERE code = ?', (code,))
    deleted_count = cursor.rowcount
    # 重新计算持仓汇总（变为空仓）
    _recalculate_portfolio(db_conn, code)
    db_conn.commit()
    audit_log(AuditAction.PORTFOLIO_REMOVE, f"清除投资记录: code={code}, deleted={deleted_count}条", extra={"code": code, "deleted_count": deleted_count})
    invalidate_portfolio_cache()
    return {'ok': True, 'deleted': deleted_count}


@router.put('/portfolio/investment/{record_id}', summary='编辑投资记录', responses={200: {'description': '更新成功'}, 400: {'description': '参数校验失败'}, 404: {'description': '记录不存在'}})
async def update_investment(record_id: int, data: InvestmentUpdate):
    """编辑投资记录（严格校验+多级兜底）"""
    db_conn = get_db()
    record = db_conn.execute('SELECT * FROM investment_records WHERE id = ?', (record_id,)).fetchone()
    if not record:
        raise HTTPException(404, '未找到该投资记录')

    # v2.8.1: 轻量路径——仅更新止盈目标（不改交易字段）时直接写库返回。
    # 此前走完整兜底校验，老记录 nav 缺失会触发"无法获取净值"400，
    # 导致前端逐条设置止盈中途失败、弹窗不关、提示失败。
    trade_fields = [data.code, data.name, data.type, data.date, data.amount, data.shares, data.nav, data.fee]
    if all(v is None for v in trade_fields) and data.profit_target is not None:
        db_conn.execute('UPDATE investment_records SET profit_target = ? WHERE id = ?',
                        (data.profit_target, record_id))
        db_conn.commit()
        invalidate_portfolio_cache()
        return {'ok': True, 'light': True}

    record_dict = dict(record)
    old_code = record_dict['code']
    old_type = record_dict['type']
    old_date = record_dict['date']
    old_amount = record_dict['amount'] or 0
    old_shares = record_dict['shares'] or 0
    old_nav = record_dict['nav'] or 0

    # ========== 参数校验 ==========
    invest_type = data.type or old_type
    valid_types = ['buy', 'sell', 'dividend', '买入', '卖出', '分红']
    if invest_type not in valid_types:
        raise HTTPException(400, f'无效的投资类型: {invest_type}，支持: buy/sell/dividend')
    type_map = {'买入': 'buy', '卖出': 'sell', '分红': 'dividend'}
    invest_type = type_map.get(invest_type, invest_type)

    new_code = data.code or old_code
    if data.code and data.code != old_code:
        fund_row = db_conn.execute('SELECT code, name, nav FROM funds WHERE code = ?', (data.code,)).fetchone()
        if not fund_row:
            raise HTTPException(404, f'基金代码 {data.code} 不存在')

    new_date = data.date or old_date
    if data.date:
        try:
            from datetime import datetime as dt
            dt.strptime(data.date, '%Y-%m-%d')
        except ValueError:
            raise HTTPException(400, f'无效的日期格式: {data.date}，请使用 YYYY-MM-DD 格式')

    new_amount = data.amount if data.amount is not None else old_amount
    new_shares = data.shares if data.shares is not None else old_shares
    if invest_type in ['buy', 'sell'] and new_amount <= 0 and new_shares <= 0:
        raise HTTPException(400, '投资金额和份额不能同时为0')

    if invest_type == 'sell' and new_shares > 0:
        portfolio = db_conn.execute('SELECT total_shares FROM portfolio WHERE code = ?', (new_code,)).fetchone()
        current_shares = portfolio['total_shares'] if portfolio else 0
        adjusted_shares = current_shares - old_shares + new_shares
        if adjusted_shares < 0:
            raise HTTPException(400, f'修改后卖出份额超过当前持有份额（当前持有{current_shares}，旧记录{old_shares}，新记录{new_shares}）')

    # ========== 净值查询兜底 ==========
    new_nav = data.nav if data.nav is not None else old_nav
    if data.date and not data.nav:
        nav_row = db_conn.execute(
            'SELECT ljjz FROM nav_history WHERE code = ? AND date = ?',
            (new_code, new_date)
        ).fetchone()
        if nav_row and nav_row[0] and nav_row[0] > 0:
            new_nav = nav_row[0]
        else:
            nav_row = db_conn.execute(
                'SELECT ljjz FROM nav_history WHERE code = ? AND date <= ? ORDER BY date DESC LIMIT 1',
                (new_code, new_date)
            ).fetchone()
            if nav_row and nav_row[0] and nav_row[0] > 0:
                new_nav = nav_row[0]
            else:
                fund_row = db_conn.execute('SELECT nav FROM funds WHERE code = ?', (new_code,)).fetchone()
                if fund_row and fund_row[0] and fund_row[0] > 0:
                    new_nav = fund_row[0]

    if not new_nav or new_nav <= 0:
        raise HTTPException(400, f'无法获取基金 {new_code} 的净值数据')

    # ========== 份额/金额自动计算 ==========
    if data.nav and not data.shares and new_amount > 0 and new_nav > 0:
        new_shares = round(new_amount / new_nav, 2)

    if data.nav and not data.amount and new_shares > 0 and new_nav > 0:
        new_amount = round(new_shares * new_nav, 2)

    if invest_type == 'sell':
        new_amount = abs(new_amount)
        new_shares = abs(new_shares)

    # ========== 基金名称兜底 ==========
    new_name = data.name
    if not new_name and (data.code or new_code != old_code):
        fund_row = db_conn.execute('SELECT name FROM funds WHERE code = ?', (new_code,)).fetchone()
        if fund_row:
            new_name = fund_row[0]
    if not new_name:
        new_name = record_dict['name']

    # ========== 构建更新字段 ==========
    update_fields = []
    update_values = []
    updates = {
        'code': new_code,
        'name': new_name,
        'type': invest_type,
        'date': new_date,
        'amount': new_amount,
        'shares': new_shares,
        'nav': new_nav,
        'fee': data.fee if data.fee is not None else record_dict['fee'],
        'profit_target': data.profit_target if data.profit_target is not None else record_dict['profit_target'],
        'note': encrypt(data.note) if data.note is not None else record_dict['note'],
    }
    for field, value in updates.items():
        update_fields.append(f'{field} = ?')
        update_values.append(value)

    update_values.append(record_id)
    db_conn.execute(f'UPDATE investment_records SET {", ".join(update_fields)} WHERE id = ?', update_values)
    db_conn.commit()

    # ========== 重新计算持仓（旧代码和新代码都要算） ==========
    _recalculate_portfolio(db_conn, old_code)
    if new_code != old_code:
        _recalculate_portfolio(db_conn, new_code)
    invalidate_portfolio_cache()

    return {
        'ok': True,
        'shares': new_shares,
        'nav': new_nav,
        'amount': new_amount,
        'type': invest_type
    }


@router.post('/portfolio/ocr-upload', summary='上传交易截图', responses={200: {'description': '上传成功，返回图片信息供手动确认'}})
async def ocr_upload_image(file: UploadFile = File(...)):
    """上传交易截图，返回图片信息供用户手动确认"""
    content = await file.read()

    upload_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'uploads')
    os.makedirs(upload_dir, exist_ok=True)

    filename = f"trade_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}_{file.filename}"
    filepath = os.path.join(upload_dir, filename)
    with open(filepath, 'wb') as f:
        f.write(content)

    return {
        'ok': True,
        'filename': filename,
        'size': len(content),
        'content_type': file.content_type,
        'image_url': f'/uploads/{filename}',
        'message': '图片已上传，请对照图片手动确认投资记录信息',
        'suggested_fields': {
            'code': '',
            'name': '',
            'type': 'buy',
            'date': datetime.date.today().isoformat(),
            'amount': 0,
            'shares': 0,
            'nav': 0,
        }
    }


@router.post('/portfolio/import/excel', summary='Excel/CSV导入投资记录', responses={200: {'description': '导入完成（含成功/跳过/错误统计）'}, 400: {'description': '文件解析失败'}})
async def import_portfolio_excel(file: UploadFile = File(...)):
    """从Excel/CSV导入投资记录（支持.csv和.xlsx格式，严格校验+多级兜底）"""
    content = await file.read()
    filename = file.filename or ''

    rows = []
    if filename.lower().endswith('.xlsx') or filename.lower().endswith('.xls'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = {}
                for i, val in enumerate(row):
                    if i < len(headers) and headers[i]:
                        row_dict[headers[i]] = val
                rows.append(row_dict)
        except Exception as e:
            raise HTTPException(400, f'Excel文件解析失败: {str(e)}')
    else:
        try:
            text = content.decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(text))
            rows = list(reader)
        except Exception as e:
            raise HTTPException(400, f'CSV文件解析失败: {str(e)}')

    db_conn = get_db()
    imported = 0
    skipped = 0
    errors = []

    for i, row in enumerate(rows, 1):
        try:
            # 表头归一化：去掉括号及括号内说明（如"交易日期（必填）"→"交易日期"，"交易类型（买入/卖出/分红）"→"交易类型"）
            normalized_row = {}
            for k, v in row.items():
                if k:
                    nk = re.sub(r'[（(][^）)]*[）)]', '', str(k)).strip()
                    normalized_row[nk] = v
                else:
                    normalized_row[k] = v
            row = normalized_row

            code = str(row.get('code') or row.get('基金代码') or row.get('基金代码/基金名称') or '').strip()
            name = str(row.get('name') or row.get('基金名称') or '').strip()
            date = str(row.get('date') or row.get('日期') or row.get('交易日期') or row.get('交易时间') or '').strip()
            type_val = str(row.get('type') or row.get('类型') or row.get('交易类型') or row.get('操作类型') or 'buy').strip().lower()
            amount_val = row.get('amount') or row.get('金额') or row.get('投资金额') or row.get('交易金额') or row.get('成交金额') or 0
            shares_val = row.get('shares') or row.get('份额') or row.get('确认份额') or row.get('成交份额') or 0
            nav_val = row.get('nav') or row.get('净值') or row.get('成交净值') or row.get('单位净值') or 0
            fee_val = row.get('fee') or row.get('手续费') or row.get('佣金') or 0
            note_val = str(row.get('note') or row.get('备注') or row.get('说明') or '').strip()

            if not code and not name:
                errors.append(f'第{i}行: 基金代码和基金名称不能同时为空')
                skipped += 1
                continue

            if not code and name:
                fund_row = db_conn.execute(
                    'SELECT code FROM funds WHERE name = ? OR name LIKE ? LIMIT 1',
                    (name, f'%{name}%')
                ).fetchone()
                if fund_row:
                    code = fund_row[0]
                else:
                    errors.append(f'第{i}行: 未找到基金名称为"{name}"的基金')
                    skipped += 1
                    continue

            fund_row = db_conn.execute('SELECT code, name, nav FROM funds WHERE code = ?', (code,)).fetchone()
            if not fund_row:
                # 基金不在行情库，自动创建一条记录（name 从 Excel 取，nav 暂时为空）
                fund_name = name or code
                db_conn.execute(
                    'INSERT OR IGNORE INTO funds (code, name, nav, nav_date) VALUES (?, ?, NULL, NULL)',
                    (code, fund_name)
                )
                db_conn.commit()
                fund_row = db_conn.execute('SELECT code, name, nav FROM funds WHERE code = ?', (code,)).fetchone()

            # 基金 nav 为空时，实时抓取最新净值（不等一键更新）
            if not fund_row or not fund_row['nav']:
                try:
                    src = EastMoneySource()
                    nav_point = src.get_latest_nav(code)
                    if nav_point and nav_point.ljjz and nav_point.ljjz > 0:
                        db_conn.execute(
                            'UPDATE funds SET nav = ?, nav_date = ? WHERE code = ?',
                            (nav_point.ljjz, nav_point.date, code)
                        )
                        db_conn.commit()
                        fund_row = db_conn.execute('SELECT code, name, nav FROM funds WHERE code = ?', (code,)).fetchone()
                except Exception:
                    pass  # 实时抓取失败不阻断导入，后续一键更新补全

            if not date:
                errors.append(f'第{i}行: 投资日期不能为空')
                skipped += 1
                continue
            try:
                from datetime import datetime as dt
                for fmt in ['%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%Y%m%d']:
                    try:
                        date = dt.strptime(date, fmt).strftime('%Y-%m-%d')
                        break
                    except ValueError:
                        continue
                else:
                    raise ValueError(f'无法解析日期格式: {date}')
            except ValueError as e:
                errors.append(f'第{i}行: 无效的日期格式 "{date}"，请使用 YYYY-MM-DD 格式')
                skipped += 1
                continue

            type_map = {'买入': 'buy', '卖出': 'sell', '分红': 'dividend', 'buy': 'buy', 'sell': 'sell', 'dividend': 'dividend'}
            invest_type = type_map.get(type_val, 'buy')

            try:
                amount = float(amount_val) if amount_val else 0
            except (ValueError, TypeError):
                amount = 0
            try:
                shares = float(shares_val) if shares_val else 0
            except (ValueError, TypeError):
                shares = 0

            if invest_type in ['buy', 'sell'] and amount <= 0 and shares <= 0:
                errors.append(f'第{i}行: 投资金额和份额不能同时为0')
                skipped += 1
                continue

            try:
                nav = float(nav_val) if nav_val else 0
            except (ValueError, TypeError):
                nav = 0

            if not nav or nav <= 0:
                nav_row = db_conn.execute(
                    'SELECT ljjz FROM nav_history WHERE code = ? AND date = ?',
                    (code, date)
                ).fetchone()
                if nav_row and nav_row[0] and nav_row[0] > 0:
                    nav = nav_row[0]

            if not nav or nav <= 0:
                nav_row = db_conn.execute(
                    'SELECT ljjz FROM nav_history WHERE code = ? AND date <= ? ORDER BY date DESC LIMIT 1',
                    (code, date)
                ).fetchone()
                if nav_row and nav_row[0] and nav_row[0] > 0:
                    nav = nav_row[0]

            if not nav or nav <= 0:
                if fund_row and fund_row['nav'] and fund_row['nav'] > 0:
                    nav = fund_row['nav']

            if not nav or nav <= 0:
                # 有投资金额但无净值时，允许导入（nav 暂为 0，后续一键更新补全）
                # 只有投资金额也为 0 时才报错
                if not amount or amount <= 0:
                    errors.append(f'第{i}行: 无法获取基金 {code} 的净值数据，且投资金额为空')
                    skipped += 1
                    continue
                # nav 保持 0，shares 后续补全

            if (not shares or shares <= 0) and amount > 0 and nav > 0:
                shares = round(amount / nav, 2)
            if (not amount or amount <= 0) and shares > 0 and nav > 0:
                amount = round(shares * nav, 2)

            if invest_type == 'sell':
                amount = abs(amount)
                shares = abs(shares)

            if not name:
                name = fund_row['name'] or code

            try:
                fee = float(fee_val) if fee_val else 0
            except (ValueError, TypeError):
                fee = 0

            db_conn.execute('''
                INSERT INTO investment_records
                (code, name, type, date, amount, shares, nav, fee, note, source)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'excel')
            ''', (code, name, invest_type, date, amount, shares, nav, fee, encrypt(note_val) if note_val else None))
            _recalculate_portfolio(db_conn, code)
            imported += 1
        except Exception as e:
            errors.append(f'第{i}行: {str(e)}')
            skipped += 1

    db_conn.commit()
    # 清除持仓缓存，确保前端立即看到导入后的数据
    if imported > 0:
        invalidate_portfolio_cache()
    return {'ok': True, 'imported': imported, 'skipped': skipped, 'errors': errors}


# ============================================================
# 持仓收益走势API
# ============================================================

@router.get('/portfolio/profit-history', summary='获取持仓收益历史走势', responses={200: {'description': '收益历史序列（含累计收益率）'}})
async def get_portfolio_profit_history(days: int = 90):
    """获取持仓收益历史走势"""
    db_conn = get_db()
    _ensure_profit_history_table(db_conn)

    rows = db_conn.execute('''
        SELECT date, total_amount, total_profit, total_cost, profit_pct, daily_profit
        FROM portfolio_profit_history
        ORDER BY date DESC
        LIMIT ?
    ''', (days,)).fetchall()

    # 按日期升序排列
    history = [dict(r) for r in reversed(rows)]

    # 计算累计收益率（以第一天为基准）
    if history:
        base_profit = history[0]['total_profit'] or 0
        base_cost = history[0]['total_cost'] or 1
        for item in history:
            item['cumulative_pct'] = round(
                ((item['total_profit'] or 0) - base_profit) / base_cost * 100, 2
            ) if base_cost > 0 else 0

    return {
        'ok': True,
        'history': history,
        'latest': history[-1] if history else None,
    }


# ============================================================
# 内部工具函数
# ============================================================

def _update_portfolio_realtime(db_conn):
    """更新持仓的实时净值和盈亏（批量查询+条件更新+非交易时段跳过）"""
    # 非交易时段跳过（周末/15:00后净值不会变化）
    now = datetime.datetime.now()
    if now.weekday() >= 5 or now.hour >= 15:
        return

    portfolios = db_conn.execute(
        'SELECT code, total_shares, avg_cost, current_nav FROM portfolio'
    ).fetchall()
    if not portfolios:
        return

    # 批量查询所有基金的净值（1次查询替代N次）
    codes = [p['code'] for p in portfolios]
    placeholders = ','.join(['?'] * len(codes))
    funds = {row['code']: row for row in db_conn.execute(
        f'SELECT code, nav, nav_date FROM funds WHERE code IN ({placeholders})',
        codes
    ).fetchall()}

    updated = False
    for p in portfolios:
        fund = funds.get(p['code'])
        if not fund or not fund['nav'] or not p['total_shares']:
            continue

        # 只在净值变化时才UPDATE（避免无意义写操作）
        if p['current_nav'] == fund['nav']:
            continue

        current_amount = round(fund['nav'] * p['total_shares'], 2)
        total_cost = round(p['avg_cost'] * p['total_shares'], 2)
        total_profit = round(current_amount - total_cost, 2)
        profit_pct = round((total_profit / total_cost * 100) if total_cost > 0 else 0, 2)
        db_conn.execute(
            "UPDATE portfolio SET current_nav = ?, nav_date = ?, total_amount = ?, "
            "total_profit = ?, profit_pct = ?, updated_at = datetime('now', 'localtime') "
            "WHERE code = ?",
            (fund['nav'], fund['nav_date'], current_amount, total_profit, profit_pct, p['code'])
        )
        updated = True

    if updated:
        db_conn.commit()


def _recalculate_portfolio(db_conn, code):
    """重新计算单只基金的持仓汇总"""
    records = db_conn.execute('''
        SELECT type, amount, shares, fee FROM investment_records WHERE code = ?
    ''', (code,)).fetchall()

    total_shares = 0
    total_cost = 0
    total_amount = 0

    for r in records:
        if r['type'] == 'buy':
            total_shares += r['shares'] or 0
            total_cost += (r['amount'] or 0) + (r['fee'] or 0)
        elif r['type'] == 'sell':
            total_shares -= r['shares'] or 0
            total_cost -= (r['amount'] or 0) - (r['fee'] or 0)
        elif r['type'] == 'dividend':
            total_cost -= r['amount'] or 0

    avg_cost = round(total_cost / total_shares, 4) if total_shares > 0 else 0

    fund = db_conn.execute('SELECT nav, nav_date, name FROM funds WHERE code = ?', (code,)).fetchone()
    current_nav = fund['nav'] if fund else 0
    nav_date = fund['nav_date'] if fund else None
    name = fund['name'] if fund else None

    current_amount = round(current_nav * total_shares, 2) if current_nav and total_shares else 0
    total_profit = round(current_amount - total_cost, 2)
    profit_pct = round((total_profit / total_cost * 100) if total_cost > 0 else 0, 2)

    if total_shares > 0:
        db_conn.execute('''
            INSERT OR REPLACE INTO portfolio
            (code, name, total_shares, avg_cost, total_amount, total_profit,
             profit_pct, current_nav, nav_date, status, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'holding', datetime('now', 'localtime'))
        ''', (code, name, total_shares, avg_cost, current_amount, total_profit,
              profit_pct, current_nav, nav_date))
    else:
        # v2.9.27 修复：清仓基金不删除，保留记录并标记为已清仓（status='sold'）
        db_conn.execute('''
            INSERT OR REPLACE INTO portfolio
            (code, name, total_shares, avg_cost, total_amount, total_profit,
             profit_pct, current_nav, nav_date, status, updated_at)
            VALUES (?, ?, 0, ?, 0, ?, 0, ?, ?, 'sold', datetime('now', 'localtime'))
        ''', (code, name, avg_cost, total_profit, current_nav, nav_date))

    db_conn.commit()
